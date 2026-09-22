"""Extract content statistics (characters, images, tables, videos, ...) from web pages and documents."""

from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import io
import json
import os
import re
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from urllib.parse import urlparse

VIDEO_HOST_PATTERN = re.compile(
    r"(youtube\.com|youtu\.be|vimeo\.com|dailymotion\.com|wistia\.|brightcove\.|player\.twitch\.tv)",
    re.I,
)
VIDEO_FILE_PATTERN = re.compile(r"\.(mp4|webm|ogv|mov|avi|mkv|m4v|wmv|flv)(\?|#|$)", re.I)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".xlsm", ".pptx")
WORDS_PER_MINUTE = 200
PREVIEW_LENGTH = 200


@dataclass
class Options:
    """Everything the analyzers can be tuned with."""

    ocr: bool = False
    ocr_language: str = "eng"
    min_rows: int = 2
    min_columns: int = 2
    pages: tuple | None = None  # (first, last), 1-based inclusive
    skip_headers: bool = False
    unique_images: bool = False
    preview: bool = False
    render: bool = False
    timeout: int = 20
    max_bytes: int = 200 * 1024 * 1024
    retries: int = 3
    cache_dir: str | None = None
    header_margin: float = 0.06


def _stats(**kwargs) -> dict:
    base = {
        "source": None,
        "type": None,
        "characters": 0,
        "characters_no_spaces": 0,
        "words": 0,
        "sentences": 0,
        "reading_time_minutes": 0.0,
        "images": 0,
        "tables": 0,
        "videos": 0,
        "links": 0,
        "headings": 0,
        "pages": None,
        "table_details": [],
    }
    base.update(kwargs)
    return base


def _text_metrics(text: str, preview: bool = False) -> dict:
    words = text.split()
    sentences = [s for s in re.split(r"[.!?\u2026]+(?:\s|$)", text) if s.strip()]
    metrics = {
        "characters": len(text),
        "characters_no_spaces": len(re.sub(r"\s", "", text)),
        "words": len(words),
        "sentences": len(sentences),
        "reading_time_minutes": round(len(words) / WORDS_PER_MINUTE, 1),
        "language": _detect_language(text),
    }
    if preview:
        metrics["preview"] = " ".join(words)[:PREVIEW_LENGTH]
    return metrics


def _detect_language(text: str) -> str | None:
    """Best-effort language code; returns None when langdetect isn't installed."""
    sample = text.strip()[:2000]
    if len(sample) < 40:
        return None
    try:
        from langdetect import detect
    except ImportError:
        return None
    try:
        return detect(sample)
    except Exception:
        return None


def _detect_kind(content_type: str, data: bytes) -> str:
    """Identify what a URL actually served, from its header and magic bytes."""
    by_mime = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    }
    if content_type in by_mime:
        return by_mime[content_type]

    if data[:4] == b"%PDF":
        return "pdf"
    if data[:2] == b"PK":
        names = _zip_names(data)
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
    return "html"


def _zip_names(data: bytes) -> list:
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return archive.namelist()
    except zipfile.BadZipFile:
        return []


def _cache_paths(url: str, opts: Options):
    digest = hashlib.sha256(url.encode()).hexdigest()
    base = os.path.join(opts.cache_dir, digest)
    return base + ".bin", base + ".type"


def _fetch(url: str, opts: Options) -> tuple[bytes, str]:
    """Download a URL with retries, a size cap and optional on-disk caching."""
    import requests
    from requests.adapters import HTTPAdapter
    from urllib3.util.retry import Retry

    if urlparse(url).scheme not in ("http", "https"):
        raise ValueError("Only http/https URLs are supported.")

    if opts.cache_dir:
        body_path, type_path = _cache_paths(url, opts)
        if os.path.exists(body_path):
            with open(body_path, "rb") as handle:
                data = handle.read()
            content_type = ""
            if os.path.exists(type_path):
                with open(type_path, encoding="utf-8") as handle:
                    content_type = handle.read().strip()
            return data, content_type

    session = requests.Session()
    session.mount(
        "https://",
        HTTPAdapter(
            max_retries=Retry(
                total=opts.retries,
                backoff_factor=0.5,
                status_forcelist=(429, 500, 502, 503, 504),
                allowed_methods=frozenset(["GET"]),
            )
        ),
    )

    with session.get(
        url,
        timeout=opts.timeout,
        stream=True,
        headers={"User-Agent": "content-stats/1.0"},
    ) as response:
        response.raise_for_status()
        declared = int(response.headers.get("Content-Length") or 0)
        if declared > opts.max_bytes:
            raise ValueError(f"Refusing to download {declared} bytes (cap {opts.max_bytes}).")

        chunks = []
        total = 0
        for chunk in response.iter_content(65536):
            total += len(chunk)
            if total > opts.max_bytes:
                raise ValueError(f"Download exceeded the {opts.max_bytes} byte cap.")
            chunks.append(chunk)
        content_type = response.headers.get("Content-Type", "").split(";")[0].strip().lower()

    data = b"".join(chunks)

    if opts.cache_dir:
        os.makedirs(opts.cache_dir, exist_ok=True)
        body_path, type_path = _cache_paths(url, opts)
        with open(body_path, "wb") as handle:
            handle.write(data)
        with open(type_path, "w", encoding="utf-8") as handle:
            handle.write(content_type)

    return data, content_type


def _render_html(url: str, opts: Options) -> bytes:
    """Load a page in a headless browser so JavaScript-built content is included."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError(
            "--render needs Playwright: pip install playwright && playwright install chromium"
        ) from exc

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        try:
            page = browser.new_page()
            page.goto(url, timeout=opts.timeout * 1000, wait_until="networkidle")
            return page.content().encode("utf-8")
        finally:
            browser.close()


def analyze_url(url: str, opts: Options | None = None) -> dict:
    opts = opts or Options()
    from bs4 import BeautifulSoup

    if opts.render:
        data, content_type = _render_html(url, opts), "text/html"
    else:
        data, content_type = _fetch(url, opts)

    kind = _detect_kind(content_type, data)
    if kind != "html":
        analyzer = {
            "pdf": analyze_pdf,
            "docx": analyze_docx,
            "xlsx": analyze_xlsx,
            "pptx": analyze_pptx,
        }[kind]
        result = analyzer(data, opts)
        result["source"] = url
        return result

    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    text = soup.get_text(separator=" ", strip=True)

    images = len(soup.find_all("img")) + len(soup.find_all("svg"))
    images += sum(1 for s in soup.find_all("source") if (s.get("type") or "").startswith("image/"))

    videos = len(soup.find_all("video"))
    for frame in soup.find_all(("iframe", "embed")):
        src = frame.get("src") or frame.get("data-src") or ""
        if VIDEO_HOST_PATTERN.search(src) or VIDEO_FILE_PATTERN.search(src):
            videos += 1
    for link in soup.find_all("a", href=True):
        if VIDEO_FILE_PATTERN.search(link["href"]):
            videos += 1

    table_details = []
    for index, table in enumerate(soup.find_all("table"), start=1):
        rows = table.find_all("tr")
        columns = max((len(r.find_all(["th", "td"])) for r in rows), default=0)
        if len(rows) < opts.min_rows or columns < opts.min_columns:
            continue
        table_details.append({"table": index, "rows": len(rows), "columns": columns})

    return _stats(
        source=url,
        type="url",
        images=images,
        tables=len(table_details),
        table_details=table_details,
        videos=videos,
        links=len(soup.find_all("a", href=True)),
        headings=len(soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"])),
        title=(soup.title.string.strip() if soup.title and soup.title.string else None),
        **_text_metrics(text, opts.preview),
    )


def analyze_docx(path, opts: Options | None = None) -> dict:
    """`path` may be a file path or the raw bytes of a .docx."""
    opts = opts or Options()
    import docx

    document = docx.Document(io.BytesIO(path) if isinstance(path, bytes) else path)

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)

    images = 0
    videos = 0
    for rel in document.part.related_parts.values():
        content_type = getattr(rel, "content_type", "")
        if content_type.startswith("image/"):
            images += 1
        elif content_type.startswith("video/"):
            videos += 1

    links = sum(
        1
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/hyperlink")
    )
    # Embedded online videos are stored as hyperlinks to a video host.
    videos += sum(
        1
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/hyperlink")
        and VIDEO_HOST_PATTERN.search(str(rel.target_ref))
    )

    table_details = [
        {"table": i, "rows": len(t.rows), "columns": len(t.columns)}
        for i, t in enumerate(document.tables, start=1)
    ]

    return _stats(
        source=None if isinstance(path, bytes) else path,
        type="docx",
        images=images,
        tables=len(table_details),
        table_details=table_details,
        videos=videos,
        links=links,
        headings=sum(
            1 for p in document.paragraphs if (p.style.name or "").startswith("Heading")
        ),
        paragraphs=len(document.paragraphs),
        **_text_metrics(text, opts.preview),
    )


def analyze_pptx(path, opts: Options | None = None) -> dict:
    """`path` may be a file path or the raw bytes of a .pptx."""
    opts = opts or Options()
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(path) if isinstance(path, bytes) else path)

    text_parts = []
    table_details = []
    images = 0
    videos = 0
    headings = 0

    def walk(shapes, slide_number):
        nonlocal images, videos, headings
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, slide_number)
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
                videos += 1
            if shape.has_text_frame:
                text_parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                table = shape.table
                table_details.append(
                    {
                        "slide": slide_number,
                        "rows": len(table.rows),
                        "columns": len(table.columns),
                    }
                )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        walk(slide.shapes, slide_number)
        if slide.shapes.title is not None:
            headings += 1
        if slide.has_notes_slide:
            text_parts.append(slide.notes_slide.notes_text_frame.text)

    return _stats(
        source=None if isinstance(path, bytes) else path,
        type="pptx",
        images=images,
        tables=len(table_details),
        table_details=table_details,
        videos=videos,
        headings=headings,
        pages=len(presentation.slides),
        slides=len(presentation.slides),
        **_text_metrics("\n".join(text_parts), opts.preview),
    )


def analyze_xlsx(path, opts: Options | None = None) -> dict:
    """`path` may be a file path or the raw bytes of an .xlsx/.xlsm workbook."""
    opts = opts or Options()
    import openpyxl

    workbook = openpyxl.load_workbook(
        io.BytesIO(path) if isinstance(path, bytes) else path,
        data_only=True,
        read_only=False,
    )

    text_parts = []
    table_details = []
    images = 0

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    text_parts.append(str(cell))

        named_tables = getattr(sheet, "tables", {})
        if named_tables:
            for name, ref in named_tables.items():
                table_details.append({"sheet": sheet.title, "named_table": name, "ref": ref})
        elif sheet.max_row > 1 and sheet.max_column > 1:
            # No declared Excel Table, so treat the used range as one.
            table_details.append(
                {"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column}
            )

        images += len(getattr(sheet, "_images", []))

    sheets = len(workbook.worksheets)
    workbook.close()

    return _stats(
        source=None if isinstance(path, bytes) else path,
        type="xlsx",
        images=images,
        tables=len(table_details),
        table_details=table_details,
        sheets=sheets,
        **_text_metrics(" ".join(text_parts), opts.preview),
    )


def analyze_pdf(path, opts: Options | None = None) -> dict:
    """`path` may be a file path or the raw bytes of a PDF."""
    opts = opts or Options()
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    text_parts = []
    images = 0
    image_xrefs = set()
    table_details = []
    videos = 0
    links = 0
    scanned_pages = 0
    analyzed_pages = 0

    def body_clip(page):
        """Rectangle excluding the running header and footer bands."""
        rect = page.rect
        inset = rect.height * opts.header_margin
        return pymupdf.Rect(rect.x0, rect.y0 + inset, rect.x1, rect.y1 - inset)

    def page_tables(page, page_number: int, text_strategy: bool, clip) -> list:
        try:
            # Scanned pages have no vector ruling lines, so align on text instead.
            # "lines_strict" ignores background fills, which otherwise look like table borders.
            strategy = "text" if text_strategy else "lines_strict"
            found = page.find_tables(strategy=strategy, clip=clip)
        except Exception:
            return []

        details = []
        for t in found.tables:
            rows = getattr(t, "row_count", len(t.rows))
            columns = getattr(t, "col_count", 0)
            if rows < opts.min_rows or columns < opts.min_columns:
                continue
            detail = {"page": page_number, "rows": rows, "columns": columns}
            try:
                cells = t.extract()
                header = [c for c in (cells[0] if cells else []) if c]
                detail["first_row"] = " | ".join(str(c).replace("\n", " ") for c in header)[:120]
            except Exception:
                pass
            details.append(detail)
        return details

    opened = (
        pymupdf.open(stream=path, filetype="pdf")
        if isinstance(path, bytes)
        else pymupdf.open(path)
    )
    with opened as doc:
        first, last = opts.pages or (1, doc.page_count)
        for page_number, page in enumerate(doc, start=1):
            if page_number < first or page_number > last:
                continue
            analyzed_pages += 1

            clip = body_clip(page) if opts.skip_headers else None
            page_text = page.get_text(clip=clip).strip()
            page_image_list = page.get_images(full=True)
            table_page = page
            scanned = False

            if not page_text and page_image_list:
                scanned_pages += 1
                scanned = True
                if opts.ocr:
                    # Rebuild the page with an OCR text layer so find_tables() has text to work with.
                    ocr_pdf = pymupdf.open(
                        "pdf",
                        page.get_pixmap(dpi=300).pdfocr_tobytes(language=opts.ocr_language),
                    )
                    table_page = ocr_pdf[0]
                    clip = body_clip(table_page) if opts.skip_headers else None
                    page_text = table_page.get_text(clip=clip).strip()

            text_parts.append(page_text)
            images += len(page_image_list)
            image_xrefs.update(item[0] for item in page_image_list)
            links += len(page.get_links())
            table_details += page_tables(table_page, page_number, scanned, clip)
            for annot in page.annots() or []:
                if annot.type[1] in ("Movie", "Screen", "RichMedia"):
                    videos += 1

        pages = doc.page_count
        headings = len(doc.get_toc())

    return _stats(
        source=None if isinstance(path, bytes) else path,
        type="pdf",
        images=len(image_xrefs) if opts.unique_images else images,
        tables=len(table_details),
        table_details=table_details,
        videos=videos,
        links=links,
        headings=headings,
        pages=pages,
        analyzed_pages=analyzed_pages,
        scanned_pages=scanned_pages,
        note=(
            f"{scanned_pages}/{analyzed_pages} analyzed pages have no text layer (scanned "
            "images); text and table counts for those pages need --ocr."
            if scanned_pages and not opts.ocr
            else None
        ),
        **_text_metrics("\n".join(text_parts), opts.preview),
    )


def analyze(source: str, opts: Options | None = None) -> dict:
    """Dispatch to the right analyzer based on the source (URL or file path)."""
    opts = opts or Options()

    if source.lower().startswith(("http://", "https://")):
        return analyze_url(source, opts)

    if not os.path.isfile(source):
        raise FileNotFoundError(source)

    extension = os.path.splitext(source)[1].lower()
    if extension == ".docx":
        return analyze_docx(source, opts)
    if extension in (".xlsx", ".xlsm"):
        return analyze_xlsx(source, opts)
    if extension == ".pptx":
        return analyze_pptx(source, opts)
    if extension == ".pdf":
        return analyze_pdf(source, opts)
    if extension == ".doc":
        raise ValueError("Legacy .doc is not supported. Convert it to .docx first.")
    if extension == ".xls":
        raise ValueError("Legacy .xls is not supported. Convert it to .xlsx first.")
    if extension == ".ppt":
        raise ValueError("Legacy .ppt is not supported. Convert it to .pptx first.")
    raise ValueError(f"Unsupported file type: {extension or '(none)'}")


def expand_sources(sources: list, recursive: bool = False) -> list:
    """Turn URLs, folders and glob patterns into a flat list of analyzable sources."""
    expanded = []
    for source in sources:
        if source.lower().startswith(("http://", "https://")):
            expanded.append(source)
        elif os.path.isdir(source):
            pattern = "**/*" if recursive else "*"
            expanded += [
                match
                for match in glob.glob(os.path.join(source, pattern), recursive=recursive)
                if match.lower().endswith(SUPPORTED_EXTENSIONS)
            ]
        elif any(char in source for char in "*?["):
            expanded += [
                match
                for match in glob.glob(source, recursive=recursive)
                if match.lower().endswith(SUPPORTED_EXTENSIONS)
            ]
        else:
            expanded.append(source)
    return expanded


def _parse_pages(value: str) -> tuple:
    match = re.fullmatch(r"\s*(\d+)\s*-\s*(\d+)\s*", value)
    if not match:
        raise argparse.ArgumentTypeError("Use the form FIRST-LAST, e.g. 8-95")
    first, last = int(match.group(1)), int(match.group(2))
    if first < 1 or last < first:
        raise argparse.ArgumentTypeError("FIRST must be >= 1 and <= LAST")
    return first, last


def write_csv(results: list, path: str) -> None:
    columns = [
        "source", "type", "characters", "characters_no_spaces", "words", "sentences",
        "reading_time_minutes", "language", "images", "tables", "videos", "links",
        "headings", "pages", "analyzed_pages", "scanned_pages", "sheets", "slides",
        "paragraphs", "title", "preview", "error",
    ]
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for result in results:
            writer.writerow(result)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "sources",
        nargs="+",
        help="URLs, files, folders or glob patterns (.pdf/.docx/.xlsx/.pptx)",
    )
    parser.add_argument("--json", action="store_true", help="print raw JSON output")
    parser.add_argument("--csv", metavar="FILE", help="write one row per source to a CSV file")
    parser.add_argument("--recursive", action="store_true", help="descend into subfolders")
    parser.add_argument(
        "--jobs", type=int, default=1, help="analyze this many sources in parallel (default: 1)"
    )
    parser.add_argument(
        "--tables", action="store_true", help="list each table with its row and column counts"
    )
    parser.add_argument("--preview", action="store_true", help="show the first 200 characters")
    parser.add_argument(
        "--pages", type=_parse_pages, metavar="FIRST-LAST", help="only analyze this PDF page range"
    )
    parser.add_argument(
        "--skip-headers",
        action="store_true",
        help="ignore the top and bottom page margins (running headers, footers, page numbers)",
    )
    parser.add_argument(
        "--unique-images",
        action="store_true",
        help="count distinct PDF images once instead of per page placement",
    )
    parser.add_argument(
        "--render",
        action="store_true",
        help="render web pages in a headless browser first (needs Playwright)",
    )
    parser.add_argument(
        "--cache", metavar="DIR", help="cache downloads in DIR to avoid re-fetching"
    )
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout in seconds")
    parser.add_argument(
        "--max-mb", type=int, default=200, help="refuse downloads larger than this (default: 200)"
    )
    parser.add_argument(
        "--retries", type=int, default=3, help="HTTP retry attempts on transient failures"
    )
    parser.add_argument(
        "--ocr", action="store_true", help="OCR scanned PDF pages (needs Tesseract installed)"
    )
    parser.add_argument(
        "--ocr-language",
        default="eng",
        help="Tesseract language code(s), e.g. eng or eng+swa (default: eng)",
    )
    parser.add_argument(
        "--min-rows",
        type=int,
        default=2,
        help="ignore detected tables with fewer rows (default: 2)",
    )
    parser.add_argument(
        "--min-columns",
        type=int,
        default=2,
        help="ignore detected tables with fewer columns (default: 2)",
    )
    args = parser.parse_args()

    opts = Options(
        ocr=args.ocr,
        ocr_language=args.ocr_language,
        min_rows=args.min_rows,
        min_columns=args.min_columns,
        pages=args.pages,
        skip_headers=args.skip_headers,
        unique_images=args.unique_images,
        preview=args.preview,
        render=args.render,
        timeout=args.timeout,
        max_bytes=args.max_mb * 1024 * 1024,
        retries=args.retries,
        cache_dir=args.cache,
    )

    sources = expand_sources(args.sources, args.recursive)

    def run(source: str) -> dict:
        try:
            return analyze(source, opts)
        except Exception as exc:
            return {"source": source, "error": f"{type(exc).__name__}: {exc}"}

    if args.jobs > 1 and len(sources) > 1:
        with ThreadPoolExecutor(max_workers=args.jobs) as pool:
            results = list(pool.map(run, sources))
    else:
        results = [run(source) for source in sources]

    if args.csv:
        write_csv(results, args.csv)
        print(f"Wrote {len(results)} rows to {args.csv}")

    if args.json:
        print(json.dumps(results, indent=2))
        return

    for result in results:
        print(f"\n=== {result['source']} ===")
        if "error" in result:
            print(f"  {result['error']}")
            continue
        for key, value in result.items():
            if key in ("source", "note", "table_details", "preview") or value is None:
                continue
            print(f"  {key.replace('_', ' ').capitalize():<22} {value}")
        if args.tables:
            for detail in result.get("table_details", []):
                location = (
                    detail.get("page")
                    or detail.get("slide")
                    or detail.get("sheet")
                    or detail.get("table")
                )
                print(
                    f"    table @ {location}: "
                    f"{detail.get('rows', '?')} rows x {detail.get('columns', '?')} columns"
                )
                if detail.get("first_row"):
                    print(f"        first row: {detail['first_row']}")
        if result.get("preview"):
            print(f"  Preview                {result['preview']}...")
        if result.get("note"):
            print(f"  Note                   {result['note']}")


if __name__ == "__main__":
    main()
